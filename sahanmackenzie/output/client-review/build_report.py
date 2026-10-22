from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

BASE = Path(__file__).resolve().parent
OUT = BASE / 'Sahan_Mackenzie_Website_Content_Review.docx'
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)
sec.top_margin, sec.bottom_margin = Inches(.70), Inches(.65)
sec.left_margin = sec.right_margin = Inches(.75)
sec.header_distance = sec.footer_distance = Inches(.30)

for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3']:
    style = doc.styles[name]
    style.font.name = 'Calibri'
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.widow_control = True
doc.styles['Normal'].font.size = Pt(10.5)
doc.styles['Normal'].paragraph_format.line_spacing = 1.12
doc.styles['Normal'].paragraph_format.space_after = Pt(6)
doc.styles['Title'].font.size = Pt(26)
doc.styles['Title'].paragraph_format.space_after = Pt(13)
doc.styles['Subtitle'].font.size = Pt(12)
doc.styles['Heading 1'].font.size = Pt(19)
doc.styles['Heading 1'].paragraph_format.space_after = Pt(12)
doc.styles['Heading 2'].font.size = Pt(12)
doc.styles['Heading 2'].paragraph_format.space_before = Pt(10)
doc.styles['Heading 2'].paragraph_format.space_after = Pt(4)
for name in ['Heading 1', 'Heading 2', 'Heading 3']:
    doc.styles[name].paragraph_format.keep_with_next = True

header = sec.header.paragraphs[0]
header.text = 'SAHAN MACKENZIE   |   WEBSITE CONTENT REVIEW'
header.runs[0].font.size = Pt(8)
header.paragraph_format.space_after = Pt(0)
footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('CONFIDENTIAL   •   22 OCTOBER 2026   •   ')
field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)
footer.add_run(' / ')
field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'NUMPAGES'); footer._p.append(field)
for run in footer.runs:
    run.font.size = Pt(8)

def p(text, bold_lead=False):
    para = doc.add_paragraph()
    if bold_lead and ': ' in text:
        first, rest = text.split(': ', 1)
        para.add_run(first + ': ').bold = True
        para.add_run(rest)
    else:
        para.add_run(text)
    return para

def h(text):
    doc.add_heading(text, level=2)

def page(title):
    doc.add_page_break()
    doc.add_heading(title, level=1)

def bullets(items):
    for item in items:
        para = doc.add_paragraph(item, style='List Bullet')
        para.paragraph_format.space_after = Pt(4)
        para.paragraph_format.line_spacing = 1.08

def table(headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, width in enumerate(widths):
        t.columns[i].width = Inches(width)
    for i, value in enumerate(headers):
        t.rows[0].cells[i].text = value
    for values in rows:
        cells = t.add_row().cells
        for i, value in enumerate(values):
            cells[i].text = str(value)
    for ri, row in enumerate(t.rows):
        rowPr = row._tr.get_or_add_trPr()
        cant = OxmlElement('w:cantSplit'); rowPr.append(cant)
        if ri == 0:
            repeat = OxmlElement('w:tblHeader'); rowPr.append(repeat)
        for ci, cell in enumerate(row.cells):
            cell.width = Inches(widths[ci])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcPr = cell._tc.get_or_add_tcPr()
            borders = OxmlElement('w:tcBorders')
            for edge in ['top', 'left', 'bottom', 'right']:
                e = OxmlElement('w:' + edge)
                for k, v in [('val', 'single'), ('sz', '5'), ('color', 'D9D9D9')]: e.set(qn('w:' + k), v)
                borders.append(e)
            tcPr.append(borders)
            margins = OxmlElement('w:tcMar')
            for edge in ['top', 'bottom', 'left', 'right']:
                e = OxmlElement('w:' + edge); e.set(qn('w:w'), '85'); e.set(qn('w:type'), 'dxa'); margins.append(e)
            tcPr.append(margins)
            shade = OxmlElement('w:shd')
            shade.set(qn('w:fill'), '303030' if ri == 0 else ('F3F4F5' if ri % 2 == 0 else 'FFFFFF'))
            tcPr.append(shade)
            for para in cell.paragraphs:
                para.paragraph_format.space_after = Pt(2)
                para.paragraph_format.line_spacing = 1.05
                if ci == 0: para.paragraph_format.keep_with_next = False
                for run in para.runs:
                    run.font.size = Pt(9)
                    run.font.color.rgb = RGBColor(255,255,255) if ri == 0 else RGBColor(0,0,0)
                    run.bold = ri == 0
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return t

doc.add_paragraph('Sahan Mackenzie Website Content Review', 'Title')
doc.add_paragraph('Recommendations from the completed founder questionnaire and current website', 'Subtitle')
p('Prepared for the website owner and implementation team. Confidential internal planning document. Reviewed 22 October 2026.')
h('Main recommendation')
p('Keep the current visual identity, color palette, imagery layout, and established animation language while developing a clearer content structure. The strongest improvements are a credible founder story, understandable service choices, concrete explanations of Design Minutiae, and a working enquiry journey. Add project proof when publication permissions and project records are ready.')
p('The current website already expresses the founder’s interest in proportion, harmony, bespoke furniture, and Sri Lankan craft. Its main weakness is that visitors see an attractive philosophy-led presentation before they can clearly understand the service options, the designer’s experience, or how to start a commission.')
h('What to change first')
bullets([
    'Resolve the public name and relationship between Mackenzie’s Atelier, Mackenzie’s Design Colonnade, and Design Minutiae Colonnade before changing the logo, domain, or metadata.',
    'Expand About into a founder-led professional story and restructure Services around how clients can engage the practice.',
    'Replace contact placeholders with approved details and add a purposeful enquiry form.',
    'Separate illustrative space categories from verified completed projects. Remove empty promises and unsupported proof.',
    'Keep internal commercial information and future ambitions separate from public website claims.'
])
h('Basis and limits of this review')
p('The source is the supplied September 2026 Founder Questionnaire, including all seven parts, 46 questions, the records appendix, and the additional brand-name notes at the end. The current local Next.js source was reviewed, including the extracted sections, content files, shared navigation, and the five new dedicated pages. This is a content and information-architecture review, not a change to the website or an audit of a deployed production site.')
p('Question references such as Q31 identify the questionnaire item supporting a finding. “Recommendation” and all draft copy describe proposed future work. The questionnaire is marked confidential and requires founder approval before publication. Several answers contain alternatives, examples, intentions, or unfilled fields; these are not treated as confirmed operating facts.')

page('1 Brand identity and positioning')
p('The public brand architecture needs a single decision. Q2 initially identifies Design Minutiae Colonnade as the brand, gives Mackenzie’s Design Colonnade or Mackenzie’s Atelier as practice-name alternatives, and assigns other phrases to philosophy, positioning, and editorial use. The closing notes favor Mackenzie’s Design Colonnade. The current website consistently leads with Mackenzie’s Atelier. These differences should be resolved with the founder, rather than silently choosing one interpretation.')
table(['Element', 'Current or source position', 'Recommendation'], [
    ('Public practice name', 'Website uses Mackenzie’s Atelier; questionnaire contains alternatives.', 'Retain the current name provisionally. Obtain one final public name before rebranding.'),
    ('Design Minutiae Colonnade', 'Called a brand name in Q2; later Design Minutiae is also described as a philosophy.', 'Agree its exact role. Explain “Design Minutiae” in plain language; avoid competing primary identities.'),
    ('Pillars of Aesthetic Harmony', 'Named philosophy; already appears in About.', 'Use as a supporting heading with practical explanations.'),
    ('The Art of Timeless Living', 'Positioning line; current hero headline.', 'Keep. Add a clear description of what the practice does.'),
    ('Where Every Detail Whispers Elegance', 'Secondary tagline in source.', 'Optional supporting copy. Avoid adding another slogan to the hero.'),
    ('The Niceties Within Mackenzie’s Design Colonnade', 'Editorial or collection concept.', 'Reserve for a journal series or future collection; keep navigation “Journal”.'),
], [1.45, 2.4, 2.92])
h('Positioning to communicate')
p('Lead with residential interior design shaped around the client’s way of living. Private residences, estates, penthouses, and coastal villas are the priority. Boutique hospitality is selective; private clubs and restaurants are lower-priority possibilities, not equal headline specialties. Sources: Q18, Q34–35, Q40, Q42.')
p('Use the founder’s preferred title, Interior Designer and Director, with final wording approved. Do not automatically adopt “Interior Architect” or “Interior Architecture & Design” from example branding notes. The supplied document does not independently establish qualifications, registrations, awards, or a studio founding date. Sources: Q1, Q3, Q9.')

page('2 What the seven parts establish')
table(['Questionnaire part', 'Website implication', 'Publication boundary'], [
    ('1 Identity and essentials\nQ1–7', 'Resolve names, title, location, contact information, and portfolio use.', 'Entity registration and public accounts are unconfirmed. Employment arrangements and private client details stay internal.'),
    ('2 Founder story\nQ8–13', 'Develop a professional biography about observation, design education, and career development.', 'Confirm career dates and qualification wording. Keep private-life content separate.'),
    ('3 Point of view and craft\nQ14–23', 'Explain Design Minutiae, bespoke work, material relationships, and Sri Lankan craftsmanship.', 'Use edited examples rather than repeating long philosophical answers.'),
    ('4 Working method\nQ24–30', 'Explain the eight-stage process, founder involvement, communication, and approvals.', 'Future systems, team size, response targets, and capacity are not current guarantees.'),
    ('5 Commercial and clients\nQ31–38', 'Clarify engagement options and qualifying enquiry questions.', 'This part is explicitly internal. Do not publish prices, margins, fee history, or client problems by default.'),
    ('6 Visibility and voice\nQ39–43', 'Add professional founder imagery, purposeful editorial content, and a consistent I/we voice.', 'Media appearances, published features, and reference brands are not endorsements or existing achievements.'),
    ('7 Vision\nQ44–46', 'Plan for project stories, original furniture, and selected international work.', 'Project targets, future team, showroom, library, and collections are ambitions.'),
], [1.55, 2.65, 2.57])
h('Important gaps despite the completed answers')
p('The questionnaire is rich in design thinking but still contains unfilled contact fields, unconfirmed capacity and support roles, incomplete project records, and commercial decisions that have not been reconciled. The appendix project and fee records do not provide completed, publishable case studies. Photographs and other supporting files are requested separately in Appendix E.')
p('The answers establish what the founder believes and intends. They do not establish permission to publish employer-owned visuals, an operating international portfolio, completed furniture collections, or independent verification of every professional claim. Treat the document as an internal briefing source, not ready-to-publish website copy.')

page('3 Current section audit first half')
h('Header and navigation')
p('Improve. Keep the visual treatment and primary enquiry action. The current navigation uses homepage anchors even though /about, /services, /spaces, /process, and /contact now exist. In a later content update, make the main navigation lead to those useful pages and use preview links within the homepage. Keep this separate from the completed structure-only refactor. Do not link to an empty Journal page. Source: Q44 and current site-header.tsx.')
h('Hero')
p('Keep and clarify. “The Art of Timeless Living” is aligned with Q2 and Q15. Add one plain sentence identifying an independent interior design practice based in Sri Lanka, focused on considered residential interiors. Keep one enquiry action and, once there is approved work, a secondary project link. Confirm image provenance and label illustrative imagery accurately. Do not imply that an AI image is a completed commission.')
h('About and philosophy preview')
p('Keep and expand. The existing paragraphs closely reflect Q12 and Q18. Add a short founder introduction and professional portrait, then link to a richer About page. Cover his habit of observation, the progression from education through brand and retail experience to interiors, and what he personally directs. Q8–11 provide substantially better material than a generic studio description. The current vase image can remain as supporting imagery.')
h('Statistics and counters')
p('Reduce. The current counters show 10+ years, eight process stages, eight craft traditions, and four design lenses. Q5 and Q9 describe approximately a decade of design leadership, not ten years of Mackenzie’s Atelier operations. Verify the starting date before retaining 10+. The other figures are content counts rather than independent proof of results. Replace most of this strip with useful design principles or move counts to the relevant sections. If a verified experience counter remains, keep the existing counter animation.')
h('Services preview')
p('Restructure. The seven current cards are capabilities, not clear ways to engage the practice. Q31 adds consultation and general documentation, multi-room residential direction, full interior design, detailed documentation, and visualization. Lead with engagement options, then show the existing seven capabilities underneath. Clarify scope and deliverables before publishing package names or fees. Retain the current card styling and carousel behavior where useful.')
h('Founder quotations')
p('Keep selectively or merge into Design Minutiae. The existing quotations are founder statements supported by Q14, Q16, Q19, and Q20. Label them as Sahan’s perspective. They must not look like client testimonials. A shorter selection will reduce repetition across the hero, About, craft, and journal sections. Preserve the slider for a purposeful selection rather than adding more philosophical slides.')

page('4 Current section audit second half')
h('Spaces')
p('Keep as project types, not a portfolio. The five current categories match the highest-ranked sectors in Q35. Make clear that the images illustrate the kinds of spaces considered. Preserve the existing AI-image disclosure and bring any label close to the relevant imagery. Only introduce a separate “Selected Projects” section after project records, exact authorship, image rights, and client permissions are documented. The internal #projects anchor currently identifies categories, not actual projects.')
h('Process')
p('Keep and explain more fully on /process. The eight existing steps are directly supported by Q25. Add the client’s input, the intended output, and the approval point for each stage. Explain that the proposal defines scope, site involvement, revisions, and responsibilities. Keep a shorter homepage preview. Do not promise fixed delivery periods or unlimited revisions; the questionnaire does not establish them.')
h('Sri Lankan craft')
p('Keep and deepen. Q22 supports the current material and craft list. Add one or two concrete explanations of how a material, making technique, or texture affects the interior. Use approved photographs of actual craft, samples, or work in progress when available. Maker names and collaborations need confirmation. Do not imply that every listed discipline has an existing commercial partnership. The current craft band is a background-image section, not a playable video.')
h('Frequently asked questions')
p('Keep and expand around decisions. The current questions are relevant. Add consultation versus full service, whether existing furniture can be retained, how references are used, what happens after an enquiry, and what the agreed proposal covers. Explain fees are scoped to the engagement without publishing internal prices. International enquiries may be welcomed, but international delivery arrangements remain subject to scope and feasibility. Sources: Q16, Q25, Q31–35, Q41.')
h('Journal preview')
p('Defer the visible block until articles are ready, or publish a small genuine first issue. Current cards all say “Coming soon” and link back to #blog rather than opening articles. Q14–23 provide strong source material, but a questionnaire answer is not automatically an approved article. Prefer one to three edited, signed articles with useful imagery over empty cards. Keep “The Niceties” as an optional editorial label under the clearer heading Journal.')
h('Contact and footer')
p('Complete before launch. Email, phone, and Instagram are placeholders, and the prominent footer invitation is an h3 rather than an actionable link. Add approved contact channels, a working enquiry route, and a minimal privacy notice reflecting actual handling. Facebook is requested in Q7 but unfilled. Do not assume WhatsApp is the preferred public channel just because it is used operationally in Q30.')
p('Revise the footer’s “selected projects abroad” wording if it implies an existing international track record. Q44 describes a first international commission as a future milestone. Safer draft direction: “Based in Sri Lanka and open to selected international enquiries.” Obtain founder approval before using it.')

page('5 Recommended homepage order')
p('Use the homepage as an introduction and route into the full site. The recommendation below changes content hierarchy, not the established palette or animation style. When approved project proof is unavailable, omit that block and move directly from the introduction to the engagement options.')
table(['Order', 'Section', 'Purpose and action'], [
    ('1', 'Hero', 'Identify the practice, residential focus, and location. Primary action: Begin a conversation.'),
    ('2', 'Founder and atelier introduction', 'Introduce Sahan, relevant verified experience, and his role. Link to About.'),
    ('3', 'Selected work when ready', 'Show a few permitted projects with clear roles. Until ready, omit rather than simulate a portfolio.'),
    ('4', 'Ways to work together', 'Present consultation, broader residential direction, and full interior design. Link to Services.'),
    ('5', 'Design Minutiae', 'Explain the approach through proportion, light, color, materials, and the client’s daily life.'),
    ('6', 'Spaces considered', 'Show the priority residential categories and selective hospitality. Link to Spaces.'),
    ('7', 'How the process works', 'Summarize the journey and the first conversation. Link to Process.'),
    ('8', 'Bespoke work and craft', 'Connect furniture, joinery, and Sri Lankan making with actual design decisions.'),
    ('9', 'Journal when ready', 'Feature a small number of complete articles. Omit empty previews.'),
    ('10', 'Questions and enquiry', 'Answer the main uncertainties, then provide a clear working contact action.'),
], [.45, 2.0, 4.32])
h('Navigation recommendation')
p('At launch: About, Services, Spaces, Process, and Contact, with the existing prominent enquiry button. Add Projects and Journal to the main navigation only when their pages contain approved content. Design Minutiae and craft can initially live within About and Services instead of requiring additional top-level navigation.')
h('Homepage depth')
p('Keep About, services, process, and editorial previews concise. Put the full founder biography, detailed inclusions, project stories, and extended explanations on dedicated pages. There is no need to expose all 46 answers, every slogan, or the internal business plan on the homepage.')

page('6 About and process page briefs')
h('About page')
p('Purpose: Help a referred visitor understand who Sahan is, why his point of view is credible, and how he works. The current /about page reuses the homepage philosophy section; it needs a fuller editorial treatment when the content update is authorized.')
bullets([
    'Founder introduction with an approved professional portrait, public name, and title.',
    'A short origin story centered on observing rooms, rearranging spaces, and noticing relationships. Edit the childhood story for relevance rather than publishing a personal memoir.',
    'Career progression: AOD and Northumbria learning; Aviraté and Timex brand and visual communication experience; design leadership at Clove Republic; the independent practice. Verify dates, qualification wording, and employer references.',
    'An explanation of Design Minutiae using two or three real decisions: a furniture proportion, material junction, or lighting relationship.',
    'Design principles covering proportion, functionality, materiality, craft, and personal meaning.',
    'Founder involvement and collaboration, without implying an established team size or services that are only planned.',
    'A link to Services and a clear enquiry action.'
])
p('Source basis: Q8–12, Q17–23, Q27–29, Q39, Q42. Keep private life separate. Professional authorship can be personal without making the website a lifestyle feed.')
h('Process page')
p('Purpose: Explain what the client can expect and what decisions keep a project moving. Retain the eight stages: enquiry; understanding the client and space; concept; design development; presentation and refinement; specification and sourcing; execution coordination; styling and handover. Source: Q25.')
p('For each stage, show what the atelier does, what information or approval the client provides, and the deliverable agreed in the proposal. Examples include a written brief, concept presentation, specification, or handover information; the exact set must follow the selected service.')
p('Add a short paragraph explaining approvals, scope changes, procurement lead times, and collaboration with contractors. Distinguish design coordination from any broader construction-management responsibility. Site-visit frequency, revision allowances, timelines, and payment milestones should be defined per engagement. Q27 and Q33 leave several of these matters unresolved.')
p('Useful reassurance: Sahan directs the concept and key aesthetic decisions while technical and specialist support can vary with the project. Do not promise that he personally manages every supplier follow-up or is available at all times. Sources: Q28–30.')

page('7 Services and engagement options')
p('The most important addition is a service structure that helps a client choose an appropriate starting point. The existing capability cards should support that structure. The following grouping is a recommendation derived from Q31, not a finalized package catalogue.')
table(['Proposed option', 'Client need', 'Scope to explain'], [
    ('Design consultation and direction', 'A considered transformation of one space.', 'Space planning, a color and material direction, furniture and lighting guidance, accessories, and an agreed selection session. Clarify the level of documentation.'),
    ('Residential design direction', 'A coherent scheme across several rooms, with the client potentially arranging execution.', 'Room-by-room planning and coordinated visual direction. Confirm which drawings, selections, sessions, and implementation support are included.'),
    ('Full interior design', 'A developed interior with coordinated details and agreed implementation involvement.', 'Concept, planning, materials, furniture, lighting, detailed documentation, bespoke design, selected visualization, and agreed sourcing and coordination.'),
], [1.55, 1.85, 3.37])
h('Capabilities to retain and clarify')
bullets([
    'Concept and design direction; space planning; color, finishes, and material palettes.',
    'Furniture selection and arrangement; bespoke furniture and joinery; lighting design.',
    'Detailed interior documentation, including relevant wall, drapery, carpet, lighting, and electrical-position information. Confirm technical coordination responsibilities.',
    '3D visualization for agreed spaces. Do not imply unlimited photorealistic renders for every engagement.',
    'Specification, sourcing, procurement support, design coordination, final styling, and agreed handover support.'
])
h('Information each option needs')
p('State who it suits, included deliverables, excluded work, client responsibilities, the degree of documentation, review stages, and the next step. Clarify whether procurement means recommendations, purchasing, supply, delivery coordination, or installation. Define whether bespoke design and fabrication are separate appointments or costs. Source: Q31–33.')
p('Q31 mentions complimentary floral and linen styling, selection days, and transportation or hospitality arrangements. Confirm the conditions before making them public entitlements. “General documentation” and “detailed documentation” also need distinct definitions. Use the website to explain value, while the proposal carries the exact scope and commercial terms.')
p('Do not publish the internal prices yet. Part 5 is explicitly marked internal, and the same document describes the LKR 3.5 million figure both as project-specific and within broader package examples. Section 12 of this report lists the commercial decisions to reconcile.')

page('8 Spaces and future project proof')
h('Spaces page')
p('Purpose: Help clients recognize whether their property and brief fit the practice. Keep the priority sequence from Q35: private residences; estates and large private homes; penthouses and luxury apartments; coastal villas and holiday residences; selective boutique hotels.')
p('For each category, explain the relevant design concerns rather than repeating luxury adjectives. Examples include circulation and family routines for residences, continuity across rooms for estates, views and daylight for penthouses, and material and climate considerations for coastal homes. Present private clubs and restaurants as selectively considered enquiries if the founder wants them mentioned; do not give them equal homepage prominence.')
h('A separate project portfolio when evidence is ready')
p('Introduce /projects separately from /spaces once there are publishable records and approved visuals. A space category explains what the studio considers; a case study demonstrates work actually undertaken. The Kandy residence mentioned in Q31 is not, by itself, a complete public case study.')
table(['Case study field', 'Required content'], [
    ('Project identity', 'Approved project title; permitted general location; year; property type; current status.'),
    ('Authorship', 'Exact design role; employer or commissioning practice where relevant; collaborators; photographer credits.'),
    ('Brief and response', 'Client requirement, architectural opportunity, principal design idea, and a meaningful design decision.'),
    ('Evidence', 'Permitted photographs, plans, detail images, or project-specific renders labeled correctly.'),
    ('Publication status', 'Written permissions and allowed channels. Distinguish public website, private presentation, and confidential material.'),
], [1.4, 5.37])
h('How to launch without project photography')
p('Use approved founder photography, original material studies, sketches, prototypes, and craft/process images to demonstrate thinking. Clearly identify concept studies and illustrative images. Omit a public project gallery if it cannot yet be supported. A private presentation can contain only material specifically permitted for that private use; a password does not replace permission.')
p('Sources: Q5–6, Q22, Q35, Q39 and Appendix A, C, E. Do not invent project names, locations, client testimonials, project counts, awards, collaborators, or outcomes to fill these templates.')

page('9 Contact and enquiry journey')
p('The current contact details are placeholders. Q7 also leaves public email, phone, social handles, domain, and preferred contact method unfilled. Completing this is a launch priority. The recommendation is one clear enquiry route supported by approved direct-contact details.')
h('Suggested first enquiry form')
table(['Field', 'Recommendation'], [
    ('Name and email', 'Required, with clear labels and useful validation.'),
    ('Phone and contact preference', 'Optional unless the founder confirms a genuine need. Add WhatsApp only when approved as a public channel.'),
    ('Property location', 'Country and city or general area. Avoid requesting a private street address initially.'),
    ('Property and project type', 'Residence, apartment, villa, selective hospitality, or another brief.'),
    ('Service interest', 'Consultation, broader residential direction, full design, or unsure.'),
    ('Scope', 'Rooms or areas, current property stage, and a short description of the desired outcome.'),
    ('Budget', 'Ask about the overall interior implementation budget; distinguish it from the design fee. Permit “not yet defined”.'),
    ('Timing', 'Desired start and any important target dates, without promising that the dates can be met.'),
    ('Referral source', 'Optional. Useful for understanding referral, social, architect, and other enquiries.'),
], [1.65, 5.12])
h('After submission')
p('Send a clear acknowledgement that confirms receipt and explains the next step: an initial review followed by a conversation if the brief appears suitable. Q30 describes response aspirations, not a confirmed public service-level promise. Do not automatically advertise a 24-hour guarantee. Route enquiries to an approved mailbox or internal workflow rather than leaving the form as a visual demo.')
h('Handling the information')
p('Add a privacy notice matching the real collection, access, and retention process. Keep marketing opt-in separate and optional if a mailing list is introduced. Start without file uploads unless the team has a clear need and a secure handling process. A later brief can collect drawings, photographs, decision-makers, and more detailed requirements.')
p('Make the footer invitation an actual link or button to /contact. Publish only approved professional channels, not private residential details. Add Instagram and Facebook links only when the accounts are ready. Sources: Q6–7, Q25, Q30, Q34, Q36–37, Q44.')

page('10 Additions removals and deferred features')
table(['Action', 'Item', 'Reason and timing'], [
    ('Add first', 'Founder profile and portrait', 'Build trust using Q8–9 and Q39. Use professional rather than private-life content.'),
    ('Add first', 'Clear engagement options', 'Consultation and full design are materially different. Q31 contains detail missing from the current site.'),
    ('Add first', 'A concrete Design Minutiae explanation', 'Show how proportion, light, materials, and daily use shape a decision. Q17–21.'),
    ('Add first', 'Working enquiry journey', 'Turn the current invitation and placeholders into a usable path. Q7 and Q44.'),
    ('Expand', 'Bespoke furniture and joinery', 'Give an important capability its own subsection within Services. Explain design, making, and scope. Q19 and Q32.'),
    ('Expand', 'Craft and material stories', 'Use real examples with context, makers and image permission where available. Q22 and Q39.'),
    ('Add when ready', 'Selected project case studies', 'The strongest proof needs completed records and publication permission. Q5–6 and Appendix A.'),
    ('Add when ready', 'Journal and founder videos', 'Publish finished content. Do not turn stated content targets into an empty website feed. Q14–23 and Q39.'),
    ('Remove or reduce', 'Decorative statistics', 'Counts of stages, crafts, and lenses do not establish project success. Retain only useful verified evidence.'),
    ('Merge', 'Repeated slogans and opinions', 'Concentrate philosophy in About and Design Minutiae; keep a shorter homepage quotation selection.'),
    ('Remove until ready', 'Coming soon journal cards', 'They currently lead back to the same section rather than to useful articles.'),
    ('Remove or qualify', 'Unproven international and team claims', 'Future commissions, studio infrastructure, and team ambitions are not current achievements.'),
    ('Defer', 'Shop, collection and showroom pages', 'Q38 and Q44–46 describe ambitions, not a ready product catalogue or operating showroom.'),
    ('Defer', 'Client portal and internal dashboards', 'Useful future operations work, but separate from the initial public marketing website.'),
], [1.0, 2.0, 3.77])
p('Also defer testimonials, press-logo strips, awards, maker-partnership badges, project filters, a newsletter, and a separate international-office section until there is real content and an operating reason. References admired in Q43 are research influences, not endorsements or partnerships.')

page('11 Voice and draft website copy')
h('Writing direction')
p('Use I for the founder’s observations, philosophy, and signed journal pieces. Use we for agreed practice services and delivery. This reflects Q42, but we should not imply a large current staff. Favor concrete explanations over repeated claims of sophistication. The source’s five preferred words are timeless, refined, layered, sophisticated, and precise; these are editorial guidance, not words to repeat in every section.')
p('Avoid unsupported superlatives and generic luxury language such as world-class, ultra-luxurious, opulent, or exclusive. Explain why a decision matters to a client. Use British English consistently, including colour, visualisation, and enquiry. The current palette and restrained visual style can remain.')
h('Draft hero supporting sentence')
p('“An independent interior design practice based in Sri Lanka, creating considered residential interiors around architecture, craftsmanship and the way you live.”')
p('Recommendation based on Q3, Q18, Q34–35. This is proposed copy, not an already approved public statement. Keep the current headline The Art of Timeless Living.')
h('Draft founder introduction')
p('“I am Sahan Mackenzie, an interior designer interested in the relationships that make a room feel complete. I begin by understanding the architecture and the person who will live within it, then consider how proportion, light, materials and furniture can work together.”')
p('Recommendation based on Q8–12 and Q17–21. Add career and experience details only after verifying the final public name, title, dates, and wording.')
h('Draft explanation of Design Minutiae')
p('“Design Minutiae is the attention given to the decisions that shape how a space feels and functions: the proportion of a table, the height of a reading light, the meeting of two materials, or the placement of something used every day. Each detail is considered in relation to the whole.”')
p('Recommendation based on Q18 and Q21. Follow it with an actual approved example rather than another slogan.')
h('Draft engagement and enquiry wording')
p('Services: “Whether you need direction for one space or a developed scheme for a residence, we begin by understanding the brief and defining the level of design involvement it needs.”')
p('Enquiry: “Tell us about your property, the spaces you would like to consider and what you hope to achieve. We will review the enquiry and discuss whether the project and our approach are a good fit.”')
p('Recommendations based on Q25, Q31, Q34 and Q36. No price, response-time guarantee, or promise of project acceptance is implied. All draft copy in this report requires founder approval before publication.')

page('12 Commercial information kept internal')
p('Part 5 explicitly says its contents stay internal. The figures below are a reconciliation checklist for the owner and team, not a recommendation to publish prices. Q31 provides useful service detail but also mixes a project-specific fee discussion with package and location examples.')
table(['Item in Q31', 'Recorded figure', 'Decision needed'], [
    ('Single-space consultation', 'From LKR 300,000', 'Define inclusions, documentation, selection-day conditions, and what the starting price excludes.'),
    ('Additional space', 'LKR 150,000', 'Confirm the definition of a space and which packages permit this rate.'),
    ('Five-bedroom consultation and general documentation', 'From LKR 1,000,000', 'Confirm scope and explain why the package differs from a simple per-space calculation.'),
    ('Five-bedroom apartment in Colombo', 'LKR 1,500,000', 'Confirm whether this is a current package, an illustration, or a historical example.'),
    ('Five-bedroom house in Colombo', 'LKR 2,500,000', 'Define area, room scope, travel, documentation, and implementation involvement.'),
    ('Outside Colombo or comprehensive residence examples', 'LKR 3,500,000', 'Reconcile the Kandy project-specific fee with the broader full-home package wording.'),
], [2.4, 1.5, 2.87])
h('Before a public pricing decision')
p('Confirm whether figures are current, whether taxes are included, what counts as design fees versus implementation costs, whether travel and hospitality are included, what site coordination covers, how many revisions and renders are allowed, and whether procurement, manufacturing, delivery, and installation are separate. Do not infer these conditions from a broad phrase such as full service.')
p('The public site can explain that a scope-specific proposal follows the initial conversation. If the founder later chooses transparent starting prices, approve a consistent commercial schedule and disclosure wording first. The source does not establish a universal minimum budget or fee, and Q34 still marks those fields to be confirmed.')
h('Other information to keep out of public sections')
p('Keep procurement margins, fee history, payment disputes, client red flags, profitability issues, staff decision rights, exact capacity, referral-contact lists, and employment-arrangement details internal. Explain relevant client responsibilities positively through the process and proposal. The questionnaire’s targets of 3–5 significant projects, a first international commission, a production-ready furniture piece, and a future team of approximately 8–12 are planning ambitions, not present website proof. Sources: Q27–38, Q44–46.')

page('13 Images permissions and credibility')
h('Use the evidence that is actually available')
p('Q5 states that Clove Republic owns company-commissioned photography, renders, and documentation, while Sahan describes substantial design responsibility over approximately a decade. His professional contribution and permission to reproduce company-owned material are separate questions. Do not present employer commissions as work independently commissioned from Mackenzie’s Atelier.')
p('For each proposed image or project, record ownership, permission status, Sahan’s exact role, employer or practice attribution, photographer or maker credit, client confidentiality conditions, and the allowed publication channels. Obtain the permissions described in Q5–6 before publishing. This report does not determine rights or interpret contracts.')
h('Image categories should remain clear')
table(['Category', 'Presentation'], [
    ('Completed commissioned project', 'Approved real project photography, accurate authorship, appropriate credits and permitted location detail.'),
    ('Project-specific render', 'Identify as a visualisation or proposed design, including project status where publishable.'),
    ('Independent concept or material study', 'Explain that it demonstrates design thinking rather than a completed commission.'),
    ('AI or other illustrative image', 'Label its illustrative purpose close to the image. Do not use it as proof of built work.'),
    ('Founder or craft photograph', 'Use approved professional imagery and identify people or makers only when cleared.'),
], [2.0, 4.77])
h('Priority asset request')
bullets([
    'One primary founder portrait and a small set of authentic working photographs.',
    'Verified biography, career dates, qualification wording, and approved professional title.',
    'A small group of candidate projects with completed Appendix A records and permission status.',
    'Original details, material boards, sketches, furniture or craft photographs that can be used publicly.',
    'Approved logo/name, professional contact details, social accounts, and domain information.',
    'Confirmed service inclusions, exclusions, proposal process, and enquiry recipient.'
])
p('Avoid invented testimonials and awards, borrowed reference imagery, employer or supplier logos without agreement, or claims of a built showroom, collection, or international office. Sources: Q5–7, Q9, Q38–39, Q43–46 and Appendix E.')

page('14 Implementation map and launch sequence')
p('The existing folder refactor is a sound base. Content improvements should now be made as a separate, reviewable phase. No website files have been changed as part of this report.')
table(['Area', 'Existing location', 'Proposed work'], [
    ('Homepage', 'src/app/page.tsx', 'Adjust section order and previews after content approval; keep current styling and behavior.'),
    ('Founder and approach', 'sections/about and content/about.ts', 'Add professional story, approved portrait, Design Minutiae examples, and fuller /about content.'),
    ('Services', 'sections/services and content/services.ts', 'Add engagement options above the existing capability cards; give /services useful depth.'),
    ('Spaces and process', 'sections/spaces, sections/process and their content files', 'Differentiate categories from projects; expand stage responsibilities and outputs.'),
    ('Contact', 'sections/contact and content/site.ts', 'Complete channels and add a functioning enquiry flow with its real handling arrangements.'),
    ('Navigation', 'components/layout', 'Link useful dedicated pages when approved; keep the primary enquiry action prominent.'),
    ('Future editorial and proof', 'New journal or projects routes when ready', 'Publish real articles and permitted case studies; defer empty dynamic routes.'),
], [1.3, 2.55, 2.92])
h('Phase one content decisions and launch essentials')
p('Confirm brand architecture and title; finalize public contacts; agree the first service structure; approve founder copy and image use; remove empty placeholders and unsupported claims; make the enquiry path work. Keep the current copper, walnut, sand, ivory, and charcoal palette and established animation behavior.')
h('Phase two useful dedicated pages')
p('Expand About, Services, Spaces, Process, and Contact beyond reused homepage sections. Add clear links between previews and full pages. Retain existing animation timings and responsive behavior unless a separate change is agreed. Check keyboard navigation, mobile legibility, reduced motion, form labels and error handling, and correct titles and descriptions.')
h('Phase three evidence and editorial')
p('Add permitted case studies, signed journal articles, bespoke object stories, and purposeful founder videos. Measure qualified enquiries and useful referral paths. Furniture collections, a showroom, a client portal, and operational dashboards belong to later scoped projects, not the initial public-site launch.')

page('15 Decisions and launch acceptance')
h('Questions for the founder')
bullets([
    'Which single public practice name should appear in the logo, domain, navigation, metadata, proposals, and social accounts? What is the exact role of Design Minutiae Colonnade?',
    'Is Interior Designer and Director the final title? Which career dates, education details, and employer references are approved for public use?',
    'Which professional email, phone, social accounts, and domain are ready? Is WhatsApp a public enquiry channel or an internal project channel?',
    'Which engagement options are currently available? What exactly separates consultation, general documentation, detailed design, and full service?',
    'Should all prices remain private? If any are published, which approved schedule and inclusions replace the mixed examples in Q31?',
    'Which projects and assets may be published, with what role attribution, permissions, credits, and location detail?',
    'Is the first international commission still a future goal? What delivery arrangements can honestly be offered to overseas enquiries today?',
    'Who reviews enquiries, and what acknowledgement and follow-up can the team reliably maintain?',
    'Which service scope, revision allowances, site involvement, procurement responsibilities, and exclusions should be stated before appointment?',
    'Can one approved article, one genuine Design Minutiae example, and a professional portrait be supplied for the first content update?'
])
h('Launch acceptance checklist')
p('A visitor can identify the designer, understand the practice’s residential focus, distinguish the engagement options, see how the process begins, and submit a real enquiry. The brand name is consistent. Contact details and links work. Every image has an accurate category and cleared use. Career and project claims are supported. Internal pricing and operational records are not exposed. Empty journal cards and unready features are omitted. The existing visual identity and animations are regression-checked after implementation.')
h('Source record')
p('Primary source: Sahan_Mackenzie_Founder_Questionnaire compleated all 7 stages.docx, supplied by the owner, marked September 2026. References in this report use its question numbers and appendix letters, including the closing brand-name notes. No external market research or independent verification of the founder’s claims was performed.')
p('Current implementation reviewed: src/app/page.tsx and layout.tsx; the /about, /services, /spaces, /process, and /contact routes; components/layout; components/sections; content/about.ts, services.ts, spaces.ts, process.ts, home.ts, journal.ts, and site.ts. Findings describe the local source at the review date, not a public deployment.')

# Remove the default Word template's decorative title border and theme residue.
for element in [doc.styles.element, doc.element]:
    for border in list(element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
for name in ['Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3']:
    style = doc.styles[name]
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.font.underline = False
    rpr = style.element.get_or_add_rPr()
    lang = OxmlElement('w:lang'); lang.set(qn('w:val'), 'en-GB'); rpr.append(lang)

replacements = {
    'color': 'colour', 'colors': 'colours', 'visualization': 'visualisation',
    'visualizations': 'visualisations', 'centered': 'centred', 'finalized': 'finalised',
    'finalize': 'finalise', 'favor': 'favour', 'labeled': 'labelled', 'labeling': 'labelling',
    'personalized': 'personalised', 'behavior': 'behaviour', 'organize': 'organise',
}
paragraphs = list(doc.paragraphs)
for t in doc.tables:
    for row in t.rows:
        for cell in row.cells:
            paragraphs.extend(cell.paragraphs)
for para in paragraphs:
    for run in para.runs:
        original_text = run.text
        revised_text = original_text
        for old, new in replacements.items():
            revised_text = re.sub(r'\b' + old + r'\b', lambda m: new.capitalize() if m.group(0)[0].isupper() else new, revised_text, flags=re.I)
        if revised_text != original_text:
            run.text = revised_text

doc.core_properties.title = 'Sahan Mackenzie Website Content Review'
doc.core_properties.subject = 'Website sections and content recommendations from the founder questionnaire'
doc.core_properties.author = 'Website content review'
doc.core_properties.keywords = 'confidential, website, content review, founder questionnaire'
doc.save(OUT)
print(OUT)
